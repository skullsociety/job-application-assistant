from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from careersgov.companion.server import JobStore
from linkedin.job_assistant.database import JobRepository
from linkedin.job_assistant.models import Job
from shared_tracker.migrate import migrate
from shared_tracker.exporter import export_database
from shared_tracker.tracking import update_tracking
from shared_tracker.resume_filenames import indexed_name
from shared_tracker.dashboard import render
from openpyxl import load_workbook


class SharedTrackerTests(unittest.TestCase):
    def test_tailored_resume_names_begin_with_the_job_id(self) -> None:
        self.assertEqual(
            indexed_name(42, "tailored_resume_Example_Data_Analyst_42.pdf"),
            "42_tailored_resume_Example_Data_Analyst.pdf",
        )
        self.assertEqual(
            indexed_name(42, "17_tailored_resume_Example_Data_Analyst.pdf"),
            "42_tailored_resume_Example_Data_Analyst.pdf",
        )

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.old = self.root / 'old/jobs.sqlite3'
        self.target = self.root / 'common/jobs.sqlite3'
        self.resumes = self.root / 'tailored'
        self.resumes.mkdir()
        self.cg = JobStore(self.old, self.root / 'resumes', self.resumes)
        self.capture = dict(title='Data Engineer', company='GovTech', url='https://jobs.careers.gov.sg/jobs/greenhouse/555',
                            job_description='Python SQL ETL', closing_date='2026-12-31', location=None)

    def tearDown(self):
        self.temp.cleanup()

    def test_migration_preserves_fields_ids_and_is_repeatable(self):
        old = self.cg.upsert(self.capture, analyze=False)
        self.cg.update_tracking(old['id'], {'applied': True, 'notes': '=literal note', 'follow_up_date': '2026-09-10'})
        (self.resumes / '1_tailored.pdf').write_bytes(b'%PDF-test')
        with closing(sqlite3.connect(self.old)) as db:
            db.execute("UPDATE jobs SET tailored_resume_path='1_tailored.pdf',match_score=67,missing_skills='docker'")
            db.commit()
        with JobRepository(self.target) as repo:
            li = repo.upsert(Job('Analyst', 'Example', 'https://www.linkedin.com/jobs/view/1/', source='linkedin'))
        preview = migrate(self.old, self.target)
        self.assertEqual(preview['source_records'], 1)
        result = migrate(self.old, self.target, apply=True, resume_root=self.resumes)
        self.assertEqual(result['inserted'], 1)
        self.assertTrue(Path(result['backup_folder']).is_dir())
        with JobRepository(self.target) as repo:
            jobs = repo.list()
            self.assertEqual(repo.get(li.id).source, 'linkedin')
            cg = next(job for job in jobs if job.source == 'careersgov')
            self.assertNotEqual(cg.id, old['id'])
            self.assertEqual(cg.legacy_careersgov_id, old['id'])
            self.assertEqual(cg.closing_date, '2026-12-31')
            self.assertEqual(cg.match_score, 67)
            self.assertIsNone(cg.location)
            self.assertTrue(cg.applied)
            self.assertTrue(cg.followed_up)
            self.assertEqual(cg.notes, '=literal note')
            self.assertEqual(cg.source_job_id, '555')
            self.assertTrue(cg.tailored_resume_path.startswith(f'{cg.id}_'))
            self.assertTrue((self.resumes / cg.tailored_resume_path).is_file())
        self.assertEqual(migrate(self.old, self.target, apply=True)['already_imported'], 1)
        with JobRepository(self.target) as repo:
            repo.delete(cg.id)
        migrate(self.old, self.target, apply=True)
        with JobRepository(self.target) as repo:
            self.assertEqual(len(repo.list()), 1)

    def test_all_adapters_read_same_schema_and_export_sources(self):
        with JobRepository(self.target) as repo:
            li = repo.upsert(Job('Analyst', 'Example', 'https://www.linkedin.com/jobs/view/1/', platform='LinkedIn', job_description='Python', key_skills='Python'))
            repo.upsert(Job('Developer', 'Example', 'https://sg.jobstreet.com/job/2', platform='JobStreet', source='jobstreet', job_description='SQL'))
        cg = JobStore(self.target, self.root / 'resumes', self.resumes)
        saved = cg.upsert(self.capture, analyze=False)
        self.assertEqual(len(cg.list_jobs()), 1)
        self.assertEqual(len(cg.list_jobs(include_all=True)), 3)
        with JobRepository(self.target) as repo:
            self.assertEqual(repo.get(saved['id']).closing_date, '2026-12-31')
            update_tracking(repo.connection, saved['id'], {'applied': True, 'notes': '=1+1'})
        output = export_database(self.target, self.root / 'jobs.xlsx')
        book = load_workbook(output)
        self.assertEqual({row[1].value for row in list(book['Job Tracker'].rows)[3:]}, {'Careers@Gov', 'LinkedIn', 'JobStreet'})
        self.assertEqual(book['Job Tracker'].max_row, 6)
        self.assertIn('Previous CareersGov ID', [cell.value for cell in book['Job Details'][1]])
        for row in list(book['Job Tracker'].rows)[3:]:
            if row[10].value == '=1+1': self.assertEqual(row[10].data_type, 's')
        book.close()
        cg.delete(saved['id'])
        export_database(self.target, output)
        book = load_workbook(output)
        self.assertEqual(book['Job Tracker'].max_row, 5)
        book.close()
        # Careers@Gov cannot rescore a LinkedIn record even when explicitly queued.
        cg.analyze_job(li.id)
        with JobRepository(self.target) as repo:
            self.assertIsNone(repo.get(li.id).match_reason)

    def test_recapture_preserves_user_notes_and_tracking(self):
        with JobRepository(self.target) as repo:
            first = repo.upsert(Job('Analyst', 'Example', 'https://www.linkedin.com/jobs/view/1/', notes='original'))
            update_tracking(repo.connection, first.id, {'notes': 'My own note', 'applied': True, 'follow_up_date': '2026-09-10'})
            second = repo.upsert(Job('Analyst', 'Example', first.url, notes='new generated description'))
            self.assertEqual(second.notes, 'My own note')
            self.assertTrue(second.applied)
            update_tracking(repo.connection, first.id, {'status': 'rejected'})
            self.assertTrue(repo.get(first.id).applied)
            self.assertEqual(repo.get(first.id).status, 'rejected')

    def test_dashboard_details_escape_captured_values(self):
        job = self.cg.upsert({**self.capture, 'company': '<script>bad</script>'}, analyze=False)
        page = render([job], 'test')
        self.assertNotIn('<script>bad</script>', page)
        self.assertIn('Closing', page)
        self.assertIn('Source Job Id', page)
        self.assertIn('Skills &amp; reasoning', page)
        self.assertNotIn('follow-up-select', page)

    def test_follow_up_date_drives_database_and_excel(self):
        with JobRepository(self.target) as repo:
            job = repo.upsert(Job('Analyst', 'Example', 'https://www.linkedin.com/jobs/view/1/'))
            for value, expected in [('2026-09-10', True), ('', False), (None, False)]:
                update_tracking(repo.connection, job.id, {'follow_up_date': value})
                self.assertEqual(repo.get(job.id).followed_up, expected)
                output = export_database(self.target, self.root / 'dates.xlsx')
                book = load_workbook(output)
                self.assertEqual(book['Job Tracker'].cell(4, 13).value, 'Yes' if expected else 'No')
                book.close()


if __name__ == '__main__':
    unittest.main()
