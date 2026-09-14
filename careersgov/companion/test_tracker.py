from __future__ import annotations

import ast
import json
import sqlite3
import tempfile
import threading
import unittest
from contextlib import closing
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from companion.excel_export import HEADERS, export_jobs
from companion.server import CompanionHandler, JobStore, render_dashboard, validate_job
from companion.test_server import SAMPLE_JOB


class TrackerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = JobStore(self.root / "jobs.sqlite3", self.root / "resumes", self.root / "tailored")

    def tearDown(self):
        self.temp.cleanup()

    def test_common_payload_and_server_owned_fields(self):
        job = self.store.upsert({
            **SAMPLE_JOB, "job_description": "Python and SQL data pipelines", "location": "Singapore",
            "salary": "SGD 5,000–7,000", "platform": "Wrong", "source": "linkedin",
            "applied": True, "match_score": 100, "status": "submitted_manually",
            "application_url": "https://example.com/apply", "workplace_type": "Hybrid",
        }, analyze=False)
        self.assertEqual(job["source"], "careersgov")
        self.assertEqual(job["platform"], "Careers@Gov")
        self.assertEqual(job["job_description"], job["description"])
        self.assertEqual(job["location"], "Singapore")
        self.assertFalse(job["applied"])
        self.assertEqual(job["status"], "saved")
        self.assertIsNone(job["match_score"])
        self.assertEqual(len(job["description_hash"]), 64)

    def test_repeated_capture_keeps_tracking_and_original_dates(self):
        first = self.store.upsert(SAMPLE_JOB, analyze=False)
        self.store.update_tracking(first["id"], {"applied": True, "followed_up": True, "notes": "Contacted recruiter", "follow_up_date": "2026-09-10"})
        second = self.store.upsert({**SAMPLE_JOB, "job_description": "Updated Python SQL description"}, analyze=False)
        self.assertEqual(first["id"], second["id"])
        self.assertEqual(first["first_seen_at"], second["first_seen_at"])
        self.assertEqual(second["seen_count"], 2)
        self.assertTrue(second["applied"])
        self.assertTrue(second["followed_up"])
        self.assertEqual(second["notes"], "Contacted recruiter")
        self.assertEqual(second["follow_up_date"], "2026-09-10")

    def test_older_analysis_cannot_overwrite_a_new_capture(self):
        old = self.store.upsert(SAMPLE_JOB, analyze=False)
        new = self.store.upsert({**SAMPLE_JOB, "job_description": "A newer SQL and Python job description."}, analyze=False)
        self.store._update_analysis(old)
        self.assertEqual(self.store.get(new["id"])["recommendation"], "analysis pending")
        self.assertEqual(self.store.get(new["id"])["job_description"], new["job_description"])

    def test_additive_legacy_migration_preserves_ids_and_resume_links(self):
        legacy = self.root / "legacy.sqlite3"
        with closing(sqlite3.connect(legacy)) as db:
            db.execute("""CREATE TABLE jobs (id INTEGER PRIMARY KEY, url TEXT UNIQUE, title TEXT, company TEXT,
                employment_type TEXT, closing_date TEXT, description TEXT, match_score INTEGER,
                tailored_resume_path TEXT, captured_at TEXT, updated_at TEXT)""")
            db.execute("INSERT INTO jobs VALUES (99, ?, 'Analyst', 'Agency', 'Contract', '31 Dec 2026', 'Python SQL', 80, '99_resume.pdf', '2026-07-01T10:00:00', '2026-07-02T10:00:00')", (SAMPLE_JOB["url"],))
            db.commit()
        migrated = JobStore(legacy, self.root / "resumes", self.root / "tailored")
        self.assertTrue(legacy.with_name(legacy.name + ".pre-v0.9.0.bak").is_file())
        job = migrated.get(99)
        self.assertEqual(job["tailored_resume_path"], "99_resume.pdf")
        self.assertEqual(job["match_score"], 80)
        self.assertEqual(job["job_description"], "Python SQL")
        self.assertEqual(job["date_found"], "2026-07-01")
        self.assertIsNone(job["location"])
        self.assertEqual(JobStore(legacy).get(99), job)

    def test_rejects_unsafe_links_and_invalid_tracking(self):
        for value in ("javascript:alert(1)", "https://user:password@example.com"):
            with self.assertRaises(ValueError):
                validate_job({**SAMPLE_JOB, "application_url": value})
        job = self.store.upsert(SAMPLE_JOB, analyze=False)
        for payload in ({"applied": "yes"}, {"status": "fake"}, {"follow_up_date": []}, {"follow_up_date": "2026-02-31"}, {"title": "changed"}):
            with self.assertRaises(ValueError):
                self.store.update_tracking(job["id"], payload)

    def test_compatible_excel_and_formula_safety(self):
        job = self.store.upsert({**SAMPLE_JOB, "title": "=1+1"}, analyze=False)
        job = self.store.update_tracking(job["id"], {"notes": "=HYPERLINK(\"https://evil.invalid\")", "follow_up_date": "2026-09-10"})
        output = export_jobs([job, job], self.root / "tracker.xlsx")
        book = load_workbook(output)
        self.assertEqual([cell.value for cell in book["Job Tracker"][3]], HEADERS)
        self.assertEqual(book["Job Tracker"].max_row, 4)
        self.assertEqual(book["Job Tracker"]["D4"].data_type, "s")
        self.assertEqual(book["Job Tracker"]["K4"].data_type, "s")
        self.assertEqual(book["Job Tracker"]["M4"].value, "Yes")
        self.assertEqual(book["Job Tracker"]["H4"].value, None)
        self.assertEqual(book["CareersGov Details"]["A2"].value, job["id"])
        book.close()
        self.store.delete(job["id"])
        export_jobs(self.store.list_jobs(), output)
        book = load_workbook(output)
        self.assertEqual(book["Job Tracker"].max_row, 3)
        book.close()

    def test_parity_with_reference_models_and_excel_columns(self):
        workspace = Path(__file__).resolve().parents[2]
        row = self.store.upsert(SAMPLE_JOB, analyze=False)
        for project in ("linkedin", "jobstreet"):
            tree = ast.parse((workspace / project / "job_assistant" / "models.py").read_text(encoding="utf-8"))
            model = next(node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "Job")
            fields = {node.target.id for node in model.body if isinstance(node, ast.AnnAssign)}
            self.assertTrue(fields.issubset(row), fields - row.keys())
            tree = ast.parse((workspace / project / "job_assistant" / "exporter.py").read_text(encoding="utf-8"))
            header_node = next(node for node in tree.body if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "HEADERS" for t in node.targets))
            self.assertEqual(HEADERS, ast.literal_eval(header_node.value))

    def test_dashboard_escapes_content_and_exposes_shared_controls(self):
        job = self.store.upsert({**SAMPLE_JOB, "title": "<script>alert(1)</script>"}, analyze=False)
        document = render_dashboard([job])
        self.assertNotIn("<script>alert(1)", document)
        for text in ("Position", "Captured", "Match", "Applied?", "Followed up?", "Resume", "Closing", "Key skills"):
            self.assertIn(text, document)

    def test_http_tracking_security_and_exports(self):
        job = self.store.upsert(SAMPLE_JOB, analyze=False)
        class Handler(CompanionHandler):
            store = self.store
            processor = None
            def log_message(self, *_):
                pass
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        connection = HTTPConnection("127.0.0.1", server.server_port)
        try:
            with patch("companion.server._refresh_excel_tracker", return_value=None) as refresh:
                path = f"/api/jobs/{job['id']}/tracking"
                for origin, expected in (("https://evil.invalid", 403), ("http://127.0.0.1:8765", 200)):
                    connection.request("POST", path, json.dumps({"applied": True}), {"Origin": origin, "Content-Type": "application/json"})
                    response = connection.getresponse()
                    self.assertEqual(response.status, expected)
                    response.read()
                refresh.assert_called_once()
            for path in ("/", "/dashboard.css", "/dashboard.js", "/api/jobs", f"/api/jobs/{job['id']}"):
                connection.request("GET", path)
                response = connection.getresponse()
                self.assertEqual(response.status, 302 if path == "/" else 200)
                if path == "/":
                    self.assertEqual(response.getheader("Location"), "http://127.0.0.1:8767/")
                response.read()
            connection.request("GET", "/dashboard/../../extension/private-profile.local.json")
            response = connection.getresponse()
            self.assertEqual(response.status, 404)
            response.read()
        finally:
            connection.close()
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == "__main__":
    unittest.main()
