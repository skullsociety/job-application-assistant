"""Local resume reading and truthful ATS-readable PDF tailoring."""

from __future__ import annotations

import re
from pathlib import Path

from .matching import MatchResult

SUPPORTED_RESUME_SUFFIXES = frozenset({".pdf", ".docx"})
SECTION_ORDER = (
    "Professional Summary",
    "Professional Experience",
    "Projects",
    "Education",
    "Certifications and Professional Development",
    "Core Skills",
    "Additional Information",
)
_MONTH = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
_DATE_RANGE = rf"{_MONTH}\s+\d{{4}}\s*(?:-|\u2013)\s*(?:Present|{_MONTH}\s+\d{{4}})"
_EXPERIENCE_START = r"(?:[A-Z][^.|]{1,90}\|\s*[^.]{1,90}|[A-Z][A-Za-z &()/\-]{2,120}(?:Resident|Intern|Consultant|Specialist|Engineer|Analyst|Officer|Technician|Developer))"
_EXPERIENCE_HEADER = re.compile(rf"(?P<header>{_EXPERIENCE_START}\s+{_DATE_RANGE})(?=\s+\d+[.)]\s*)")


def latest_resume(directory: Path) -> Path:
    if not directory.is_dir():
        raise FileNotFoundError(f"Resume folder not found: {directory}")
    candidates = [path for path in directory.iterdir() if path.is_file() and path.suffix.casefold() in SUPPORTED_RESUME_SUFFIXES]
    if not candidates:
        raise FileNotFoundError(f"No PDF or DOCX resume found in: {directory}")
    return max(candidates, key=lambda path: (path.stat().st_mtime_ns, path.name.casefold()))


def read_resume(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")
    if path.suffix.casefold() == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ValueError("The PDF resume is password-protected.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif path.suffix.casefold() == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        text = "\n".join(paragraphs + cells)
    else:
        raise ValueError("Resume must be a PDF or DOCX file.")
    if not text.strip():
        raise ValueError(f"No readable text was found in {path.name}.")
    return text.strip()


def create_tailored_resume(
    resume_text: str,
    job_id: int,
    company: str,
    title: str,
    match: MatchResult,
    output_dir: Path,
) -> Path:
    """Create a PDF using only source-resume content and verified matching skills."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{job_id}_tailored_resume_{_safe_name(company)}_{_safe_name(title)}.pdf"
    from uuid import uuid4
    temporary = output.with_name(f".{output.stem}.{uuid4().hex}.building.pdf")
    document = SimpleDocTemplate(
        str(temporary),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=17 * mm,
        bottomMargin=14 * mm,
        title=f"Tailored Resume - {title}",
        author="Careers@Gov Assistant",
    )
    base = getSampleStyleSheet()
    styles = {
        "name": ParagraphStyle("Name", parent=base["Title"], fontName="Helvetica-Bold", fontSize=15, leading=18, alignment=TA_CENTER, textColor=colors.HexColor("#172033"), spaceAfter=1),
        "contact": ParagraphStyle("Contact", parent=base["Normal"], fontName="Helvetica", fontSize=9.5, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#465267")),
        "heading": ParagraphStyle("Heading", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#702459"), spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "body": ParagraphStyle("Body", parent=base["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.2, spaceAfter=5),
        "experience_header": ParagraphStyle("ExperienceHeader", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9.7, leading=12.4, textColor=colors.HexColor("#34425a"), spaceBefore=5, spaceAfter=3, keepWithNext=True),
        "bullet": ParagraphStyle("Bullet", parent=base["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.2, leftIndent=12, firstLineIndent=-8, bulletIndent=0, spaceAfter=4),
    }
    contact, sections = _split_resume(resume_text)
    name, details = _contact_lines(contact)
    story = [Paragraph(_escape(name or "Resume"), styles["name"])]
    if details:
        story.append(Paragraph(_escape(details), styles["contact"]))
    story.extend([Spacer(1, 8), Paragraph("Relevant Skills", styles["heading"])])
    verified = ", ".join(_display_skill(skill) for skill in match.matching_skills)
    story.append(Paragraph(_escape(verified or "No verified matching skills were detected."), styles["body"]))
    for heading in SECTION_ORDER:
        content = sections.get(heading)
        if not content:
            continue
        story.append(Paragraph(_escape(heading), styles["heading"]))
        if heading == "Professional Experience":
            story.extend(_experience_flowables(content, styles, match.matching_skills))
        else:
            story.extend(_content_paragraphs(content, styles))
    if not sections:
        story.append(Paragraph("Original Resume Content", styles["heading"]))
        story.extend(_content_paragraphs(resume_text, styles))
    try:
        document.build(story)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def _split_resume(resume_text: str) -> tuple[str, dict[str, str]]:
    text = _normalise(resume_text)
    pattern = "|".join(re.escape(heading) for heading in sorted(SECTION_ORDER, key=len, reverse=True))
    matches = list(re.finditer(pattern, text, re.IGNORECASE))
    if not matches:
        return text[:200], {}
    contact = text[:matches[0].start()].strip()
    sections: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        heading = next(item for item in SECTION_ORDER if item.casefold() == match.group(0).casefold())
        sections[heading] = text[match.end():end].strip()
    return contact, sections


def _contact_lines(contact: str) -> tuple[str, str]:
    phone = re.search(r"\b\d{8,}\b", contact)
    if phone:
        return contact[:phone.start()].strip(" |"), contact[phone.start():].strip()
    parts = re.split(r"\s*\|\s*", contact, maxsplit=1)
    return (parts[0], parts[1]) if len(parts) > 1 else (contact, "")


def _content_paragraphs(content: str, styles: dict[str, object]) -> list[object]:
    from reportlab.platypus import Paragraph

    chunks = [chunk.strip() for chunk in re.split(r"(?=\b\d+[.)]\s*)|\n{2,}", content) if chunk.strip()]
    output: list[object] = []
    for chunk in chunks:
        numbered = bool(re.match(r"^\d+[.)]", chunk))
        cleaned = re.sub(r"^\d+[.)]\s*(?:-\s*)?", "", chunk).strip()
        output.append(Paragraph(_escape(cleaned), styles["bullet"] if numbered else styles["body"], bulletText="-" if numbered else None))
    return output


def _experience_flowables(content: str, styles: dict[str, object], matching_skills: list[str]) -> list[object]:
    from reportlab.platypus import KeepTogether, Paragraph, Spacer

    blocks = _experience_blocks(content)
    if not blocks:
        return _content_paragraphs(content, styles)
    ranked = sorted(
        enumerate(blocks),
        key=lambda item: (-_experience_relevance(item[1], matching_skills), item[0]),
    )
    output: list[object] = []
    for _, (header, bullets) in ranked:
        heading = Paragraph(_escape(_ascii_text(header)), styles["experience_header"])
        bullet_flowables = [
            Paragraph(_escape(_ascii_text(_clean_numbered_item(item))), styles["bullet"], bulletText="-")
            for item in bullets
        ]
        if bullet_flowables:
            output.append(KeepTogether([heading, bullet_flowables[0]]))
            output.extend(bullet_flowables[1:])
        else:
            output.append(heading)
        output.append(Spacer(1, 5))
    return output


def _experience_blocks(content: str) -> list[tuple[str, list[str]]]:
    matches = list(_EXPERIENCE_HEADER.finditer(content))
    blocks: list[tuple[str, list[str]]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(content)
        body = content[match.end():end]
        bullets = [part.strip() for part in re.split(r"(?=\b\d+[.)]\s*)", body) if part.strip()]
        blocks.append((match.group("header").strip(), bullets))
    return blocks


def _experience_relevance(block: tuple[str, list[str]], matching_skills: list[str]) -> int:
    text = " ".join((block[0], *block[1])).casefold()
    return sum(skill.casefold() in text for skill in matching_skills)


def _clean_numbered_item(value: str) -> str:
    return re.sub(r"^\d+[.)]\s*(?:-\s*)?", "", value).strip()


def _normalise(text: str) -> str:
    return _ascii_text(re.sub(r"\s+", " ", text).strip())


def _ascii_text(value: str) -> str:
    replacements = {
        "\u2011": "-",
        "\u2012": "-",
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "-",
        "\u00a0": " ",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    return re.sub(r"\.\s+\.", ".", value)


def _escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")[:48] or "job"


def _display_skill(skill: str) -> str:
    upper = {"sql", "aws", "gcp", "api", "uat", "jira", "ci/cd", "html", "css"}
    return skill.upper() if skill in upper else skill.title()
